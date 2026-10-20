"""Create fictional sample files for trying every import in regTester.

    python tools/make_samples.py        -> writes into sample-data/

Needs openpyxl and python-docx only for the .xlsx/.docx samples (pip install openpyxl python-docx);
the CSV/VTT samples are written with the standard library.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "sample-data"


def write_csv(name, header, rows):
    with open(OUT / name, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(header)
        w.writerows(rows)


def main():
    OUT.mkdir(exist_ok=True)
    # ------------------------------------------- regtest tool & SharePoint exports
    # Fictional data in the exact column layout of the real exports (see tools/sample_regtool.json;
    # the browser demo uses the same data).
    data = json.loads((Path(__file__).with_name("sample_regtool.json")).read_text(encoding="utf-8"))
    try:
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.title = "Timeline"
        ws.append(["Timeline Version", "Schlagwort", "Vorgangsname", "Anfangsdatum", "Fälligkeitsdatum", "Zugewiesen an", "Anzeigen"])
        for r in data["timeline"]:
            ws.append([r["version"], r["keyword"], r["title"], r["start"], r["end"], r["assigned"], r["show"]])
        wb.save(OUT / "timeline_sharepoint_export.xlsx")

        em_head = ["Error-ID", "Title", "Error Status", "Error Type", "TestRunCount", "Affected TestPackages", "Responsibles",
                   "First Time Stamp", "Last Time Stamp", "Application Name"]
        for n, key in ((1, "errors_1"), (2, "errors_2")):
            wb = Workbook(); ws = wb.active; ws.title = "Errors"
            ws.append(["Demo Error Report"])
            ws.append(em_head)
            for e in data[key]:
                ws.append([e["id"], e["title"], e["status"], e["type"], e["runs"], e["packages"], e["responsibles"], e["first"], e["last"], e["app"]])
            wb.save(OUT / f"error_management_export_{n}.xlsx")

        wb = Workbook(); ws = wb.active; ws.title = "TestCases"
        for r in data["catalog"]:
            ws.append(r)
        wb.save(OUT / "testcases_list_export.xlsx")

        wb = Workbook(); ws = wb.active; ws.title = "Tree"
        ws.append(["TestCase Reporting Tree"])
        tree = {}
        for r in data["catalog"][2:]:
            if r[5]:
                tree.setdefault(r[5], {}).setdefault(r[6], {}).setdefault(r[7], []).append(r[2])
        tree.setdefault("Orders", {}).setdefault("Import", {}).setdefault("other TestCases", []).append("Import orders from CSV")
        for main, subs in tree.items():
            ws.append([main])
            for sub, secs in subs.items():
                ws.append(["", sub])
                for sec, titles in secs.items():
                    ws.append(["", "", sec])
                    for title in titles:
                        ws.append(["", "", "", title])
        wb.save(OUT / "testcases_tree_export.xlsx")

        wb = Workbook(); ws = wb.active; ws.title = "Documents"
        ws.append(["Demo Document Report"])
        ws.append(["No", "Title", "ID", "Milestone", "Idents", "Status", "Created", "CreatedBy", "Related Items", "CheckedOut", "Hint"])
        for i, d in enumerate(data["docs"], 1):
            ws.append([i, d["title"], d["id"], d["milestone"], d["idents"], d["status"], d["created"], "demo-user", d["related"], "", ""])
        wb.save(OUT / "document_management_export.xlsx")

        wb = Workbook(); ws = wb.active; ws.title = "Bugs"
        ws.append(["Issue key", "Issue Type", "Summary", "Description", "Status", "Component", "Fix Version", "Labels"])
        ws.append(["BUG-2301", "Bug", "Export button disabled for auditor", "Auditor sees the export button greyed out although read access exists.", "Open", "Admin", "V8.3", "auditor, export"])
        ws.append(["BUG-2302", "Bug", "Cost center not copied to invoice", "When an order becomes an invoice the cost center is empty.", "In Progress", "Invoice", "V8.3", "costcenter"])
        ws.append(["ISS-95", "Issue", "Sandbox mail server not reachable", "Reminder mails cannot be tested in sandbox since Monday.", "Open", "General", "", "sandbox, mail"])
        wb.save(OUT / "bugs_export.xlsx")
    except ImportError:
        print("openpyxl missing: .xlsx samples skipped")

    # ----------------------------------------------------- release notes
    feats = [("FN-1042", "Export PDF facturi", "S39", "Invoice", "Invoices can be exported as PDF from the list."),
             ("FN-1051", "Filtru avansat clienți", "S40", "Invoice", "Filter invoices by customer region and cost center."),
             ("FN-1055", "Import comenzi din CSV", "S40", "Orders", "Orders can be imported from CSV files."),
             ("FN-1060", "Notificări e-mail scadență", "S41", "Invoice", "Reminder e-mail three days before due date.")]
    write_csv("release_notes.csv", ["ID", "Title", "Sprint", "Release", "Module", "Description"], [[k, t, s, "V8.3", m, d] for k, t, s, m, d in feats])
    try:
        from docx import Document
        doc = Document()
        doc.add_heading("Release Notes V8.3", 0)
        for k, t, s, m, d in feats:
            doc.add_heading(f"{k} {t}", level=2)
            doc.add_paragraph(f"Sprint {s} · Modul {m}")
            doc.add_paragraph(d)
        doc.save(OUT / "release_notes.docx")

        doc = Document()
        doc.add_heading("Prozeduren Regtest", 0)
        doc.add_heading("PR-130 Mahnlauf im Sandbox prüfen", level=1)
        doc.add_paragraph("Mahnlauf manuell starten, Protokoll prüfen, E-Mail im Testpostfach kontrollieren.")
        doc.add_paragraph("Kunde K-71 verwenden", style="List Bullet")
        doc.add_paragraph("Fälligkeit auf heute + 3 Tage setzen", style="List Bullet")
        doc.add_heading("PR-131 CSV-Import von Aufträgen", level=1)
        doc.add_paragraph("Datei mit 10 Zeilen importieren, Fehlerprotokoll pro Zeile prüfen.")
        t = doc.add_table(rows=3, cols=2)
        for i, (a, b) in enumerate([("Feld", "Erwartung"), ("Trennzeichen", "Semikolon oder Komma"), ("Kodierung", "UTF-8")]):
            t.cell(i, 0).text, t.cell(i, 1).text = a, b
        doc.save(OUT / "procedures_onenote_export.docx")

        doc = Document()
        doc.add_paragraph("Sprint 41 Review")
        for name, time, text in [("Ravi Kumar", "0:12", "so basically uh the in voice reminder will go three days before due date"),
                                 ("Jonas Müller", "0:25", "ähm und das geht auch für Kunden ohne E-Mail Adresse?"),
                                 ("Ravi Kumar", "0:31", "no no only customers with mail, others get the letter like before you know")]:
            doc.add_paragraph(f"{name}   {time}")
            doc.add_paragraph(text)
        doc.save(OUT / "teams_transcript_sprint41.docx")
    except ImportError:
        print("python-docx missing: .docx samples skipped")

    # ----------------------------------------------------- regtest runs
    base = [f"TC-{3000 + i}" for i in range(30)] + ["TC-1800", "TC-1801", "TC-2201", "TC-2220"]
    runs = {
        "2026-10-31": {"TC-1800": "Rounding diff 0.01 in VAT total (expected 714.00, got 714.01)", "TC-1801": "Test data missing: customer K-77",
                       "TC-3007": "Element #btnSave not clickable after 15s", "TC-3008": "Element #btnSave not clickable after 15s"},
        "2026-11-03": {"TC-1800": "Rounding diff 0.01 in VAT total (expected 99.00, got 99.01)", "TC-3007": "Element #btnSave not clickable after 15s",
                       "TC-2220": "Field 'CostCenter' is mandatory (order 90011)"},
    }
    for date, fails in runs.items():
        write_csv(f"regtest_run_{date}.csv", ["Test", "Result", "Error message", "Date"],
                  [[t, "Failed" if t in fails else "Passed", fails.get(t, ""), date] for t in base])

    (OUT / "transcript_sprint41.vtt").write_text(
        "WEBVTT\n\n00:00:01.000 --> 00:00:04.000\n<v Ravi K.>okay okay so today uh the in voice reminder</v>\n\n"
        "00:00:04.500 --> 00:00:09.000\n<v Ravi K.>it goes three days before due date for every custom are with e-mail</v>\n\n"
        "00:00:09.500 --> 00:00:12.000\n<v Jonas M.>ähm und was passiert bei Kunden ohne E-Mail?</v>\n\n"
        "00:00:12.500 --> 00:00:16.000\n<v Ravi K.>they get the letter like before, the reg test should check both</v>\n", encoding="utf-8")
    write_csv("glossary.csv", ["Heard as", "Correct term"], [["in voice", "Invoice"], ["custom are", "customer"], ["reg test", "regtest"], ["cost centre", "cost center"]])
    print("Samples written to", OUT)


if __name__ == "__main__":
    main()
