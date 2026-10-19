"""Write one blank-ish Excel template per upload in regTester, with example rows.

    python tools/make_templates.py      -> writes into sample-data/templates/

Each workbook has a single sheet (the importer picks the first sheet that has data).
Orange headers are required; hover a header in Excel to see the other names it accepts.
Needs openpyxl (pip install openpyxl).
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent.parent / "sample-data" / "templates"
REQ = PatternFill("solid", fgColor="F4B183")
OPT = PatternFill("solid", fgColor="D9E1F2")
D = dt.date

# (file, sheet, [(header, required, comment)], rows, date columns)
TEMPLATES = [
    ("1_timeline.xlsx", "Timeline", [
        ("Tip", False, "sprint, planning, review, deployment, regtest, qa, freeze, holiday, other.\n"
                       "Merg și cuvinte libere: Sprintplanung, Go-Live, Regressionstest, Q&A, Feiertag, concediu...\n"
                       "Gol = ghicit din titlu.\nAlte nume de coloană: Type, Typ, Category, Art"),
        ("Titlu", True, "Alte nume: Title, Name, Titel, Denumire, Eveniment, Event"),
        ("Început", True, "Dată Excel sau text: 2027-01-04, 04.01.2027, 04/01/2027.\nAlte nume: Start, Beginn, Von, From, Data"),
        ("Sfârșit", False, "Gol = aceeași zi cu Început.\nAlte nume: End, Ende, Bis, To, Finish"),
        ("Release", False, "Alte nume: Version, Versiune"),
        ("Sprint", False, "Alte nume: Iteration"),
        ("Responsabil", False, "Alte nume: Owner, Verantwortlich, Assigned to"),
        ("Note", False, "Alte nume: Notes, Notizen, Comment, Comentariu"),
        ("ID", False, "Opțional. Gol = se folosește titlul ca cheie (pentru «actualizează după cheie»)."),
    ], [
        ["sprint", "Sprint 60", D(2027, 1, 4), D(2027, 1, 15), "R-2027.1", "Sprint 60", "Team Blau", "", ""],
        ["planning", "Sprint 60 Planning", D(2027, 1, 4), D(2027, 1, 4), "R-2027.1", "Sprint 60", "PO", "Teams", ""],
        ["review", "Sprint 60 Review", D(2027, 1, 15), D(2027, 1, 15), "R-2027.1", "Sprint 60", "PO", "Demo", ""],
        ["freeze", "R-2027.1 Code freeze", D(2027, 2, 12), D(2027, 2, 12), "R-2027.1", "", "Dev", "", ""],
        ["regtest", "R-2027.1 Regtest", D(2027, 2, 15), D(2027, 2, 23), "R-2027.1", "", "QA", "", ""],
        ["qa", "R-2027.1 Q&A", D(2027, 2, 24), D(2027, 2, 26), "R-2027.1", "", "QA", "", ""],
        ["deployment", "R-2027.1 Go-Live", D(2027, 3, 1), D(2027, 3, 1), "R-2027.1", "", "Ops", "", ""],
        ["holiday", "Anul Nou", D(2027, 1, 1), D(2027, 1, 1), "", "", "", "", ""],
    ], {3, 4}),

    ("2_acoperire_release_notes.xlsx", "Release notes", [
        ("Cod", False, "Codul funcției, folosit pentru legătura cu testele.\nAlte nume: ID, Key, Ticket, Story, Jira, PBI, Work item"),
        ("Titlu", True, "Alte nume: Title, Name, Feature, Funktion, Functie, Summary"),
        ("Sprint", False, "Alte nume: Iteration"),
        ("Release", False, "Alte nume: Version, Versiune"),
        ("Modul", False, "Alte nume: Module, Area, Bereich, Component, Componenta"),
        ("Descriere", False, "Alte nume: Description, Beschreibung, Details, Text"),
    ], [
        ["FN-1042", "Export PDF facturi", "S39", "R-2026.4", "Invoice", "Facturile pot fi exportate ca PDF din listă."],
        ["FN-1051", "Filtru avansat clienți", "S40", "R-2026.4", "Invoice", "Filtrare după regiune și centru de cost."],
        ["FN-1055", "Import comenzi din CSV", "S40", "R-2026.4", "Orders", "Comenzile pot fi importate din fișiere CSV."],
    ], set()),

    ("3_acoperire_teste.xlsx", "Teste", [
        ("ID test", True, "Alte nume: Test ID, Testcase, TC, ID, Key"),
        ("Titlu", True, "Alte nume: Title, Name, Test case, Testfall, Denumire"),
        ("Modul", False, "Alte nume: Module, Area, Folder, Ordner, Suite"),
        ("Cod funcție", False, "Codul din release notes (ex. FN-1055) -> leagă testul de funcție.\n"
                               "Alte nume: Feature, Requirement, Anforderung, Cerinta, Story, Link"),
        ("Status", False, "Text liber (Ready, Draft...).\nAlte nume: State, Zustand, Stare"),
    ], [
        ["TC-2201", "Import comenzi CSV - happy path", "Orders", "FN-1055", "Ready"],
        ["TC-2202", "Import comenzi CSV - rând invalid", "Orders", "FN-1055", "Draft"],
        ["TC-2210", "Export PDF factură", "Invoice", "FN-1042", "Ready"],
    ], set()),

    ("4_triere_rezultate_regtest.xlsx", "Rulare", [
        ("Test", True, "ID-ul testului, la fel ca în exportul de teste.\nAlte nume: Test ID, Testcase, TC, ID, Name, Testfall"),
        ("Rezultat", True, "Failed / Passed. Se recunosc și: Error, Fehler, NOK, Esuat, OK, Success, Bestanden, Reusit.\n"
                           "Orice altceva = Skipped.\nAlte nume: Status, Result, Ergebnis, Outcome, Verdict"),
        ("Mesaj de eroare", False, "Erorile se grupează după mesaj (cifrele și datele sunt ignorate).\n"
                                   "Alte nume: Message, Error, Fehlermeldung, Eroare, Mesaj, Reason"),
        ("Data", False, "Data rulării; gol = azi. Un fișier = o rulare.\nAlte nume: Date, Datum, Run date, Timestamp"),
    ], [
        ["TC-2201", "Passed", "", D(2026, 11, 3)],
        ["TC-2202", "Failed", "Field 'CostCenter' is mandatory (order 90011)", D(2026, 11, 3)],
        ["TC-2210", "Failed", "Rounding diff 0.01 in VAT total (expected 99.00, got 99.01)", D(2026, 11, 3)],
        ["TC-3007", "Failed", "Element #btnSave not clickable after 15s", D(2026, 11, 3)],
        ["TC-3008", "Skipped", "", D(2026, 11, 3)],
    ], {4}),

    ("5_triere_erori_cunoscute.xlsx", "Erori cunoscute", [
        ("Tipar", True, "Bucată din mesajul de eroare; dacă mesajul o conține, eroarea e marcată «cunoscută».\n"
                        "Alte nume: Pattern, Message, Error, Fehler, Eroare"),
        ("Motiv", False, "Alte nume: Reason, Grund, Comment"),
        ("Bug / tichet", False, "Alte nume: Bug, Ticket, Issue, Defect"),
    ], [
        ["not clickable after", "Instabilitate UI în sandbox", "BUG-2301"],
        ["rounding diff", "Rotunjire TVA, fix în R-2027.1", "BUG-2302"],
    ], set()),

    ("6_specificatii.xlsx", "Specificatii", [
        ("ID", False, "Alte nume: Key, Nr, Number, Ticket, Bug ID, Issue ID, Cod"),
        ("Titlu", True, "Alte nume: Title, Summary, Zusammenfassung, Subject, Betreff, Name"),
        ("Tip", False, "bug, issue, procedure, spec, note (merg și Defect, Fehler, Incident, Anleitung, Requirement, Cerinta...).\n"
                       "Gol = tipul ales la butonul de import.\nAlte nume: Type, Typ, Category, Issue type"),
        ("Text", False, "Alte nume: Description, Beschreibung, Descriere, Details, Content, Steps"),
        ("Status", False, "Alte nume: State, Zustand, Stare"),
        ("Modul", False, "Alte nume: Module, Component, Komponente, Area"),
        ("Release", False, "Alte nume: Version, Fix version, Versiune"),
        ("Etichete", False, "Separate prin virgulă sau ;\nAlte nume: Tags, Labels, Schlagworte"),
    ], [
        ["BUG-2301", "Buton export dezactivat pentru auditor", "bug", "Auditorul vede butonul gri deși are drept de citire.", "Open", "Admin", "R-2026.4", "auditor, export"],
        ["ISS-95", "Server mail sandbox indisponibil", "issue", "Mailurile de reamintire nu pot fi testate.", "Open", "General", "", "sandbox, mail"],
        ["PR-130", "Mahnlauf în sandbox", "procedure", "1. Pornește manual\n2. Verifică protocolul\n3. Verifică mailul", "", "Invoice", "", "mahnung"],
        ["SPEC-12", "Centru de cost obligatoriu pe comandă", "spec", "Comanda nu se salvează fără centru de cost.", "Approved", "Orders", "R-2026.4", ""],
    ], set()),

    ("7_sesiuni_po_glosar.xlsx", "Glosar", [
        ("Auzit ca", True, "Cum apare greșit în transcript.\nAlte nume: Wrong, Falsch, Gresit, Heard, Recognized"),
        ("Termen corect", True, "Alte nume: Right, Correct, Richtig, Corect, Term"),
    ], [
        ["in voice", "Invoice"],
        ["custom are", "customer"],
        ["reg test", "regtest"],
        ["cost centre", "cost center"],
    ], set()),
]


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for name, sheet, cols, rows, date_cols in TEMPLATES:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet
        ws.append([c[0] for c in cols])
        for i, (header, required, note) in enumerate(cols, 1):
            cell = ws.cell(1, i)
            cell.font = Font(bold=True)
            cell.fill = REQ if required else OPT
            cell.comment = Comment(("OBLIGATORIU\n" if required else "") + note, "regTester", width=320, height=140)
            width = max(len(header), *(len(str(r[i - 1]).split("\n")[0]) for r in rows)) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(max(width, 12), 60)
        for r in rows:
            ws.append(r)
        for c in date_cols:
            for cell in ws[get_column_letter(c)][1:]:
                cell.number_format = "DD.MM.YYYY"
        ws.freeze_panes = "A2"
        wb.save(OUT / name)
    print("Templates written to", OUT)


if __name__ == "__main__":
    main()
